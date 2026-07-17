import fitz
from docx import Document
from pptx import Presentation


def read_pdf(file_bytes: bytes) -> str:
    text = ""
    with fitz.open(stream=file_bytes, filetype="pdf") as doc:
        for page in doc:
            text += page.get_text()
    return text


def read_docx(file_bytes: bytes) -> str:
    import io
    doc = Document(io.BytesIO(file_bytes))
    return " ".join([p.text for p in doc.paragraphs])


def read_pptx(file_bytes: bytes) -> str:
    import io
    prs = Presentation(io.BytesIO(file_bytes))
    return " ".join(
        shape.text
        for slide in prs.slides
        for shape in slide.shapes
        if hasattr(shape, "text")
    )


def read_txt(file_bytes: bytes) -> str:
    return file_bytes.decode("utf-8")


def extract_text(filename: str, file_bytes: bytes) -> str:
    ext = filename.rsplit(".", 1)[-1].lower()
    readers = {
        "pdf": read_pdf,
        "docx": read_docx,
        "pptx": read_pptx,
        "txt": read_txt,
    }
    reader = readers.get(ext)
    if not reader:
        raise ValueError(f"Unsupported file type: .{ext}")
    return reader(file_bytes)
